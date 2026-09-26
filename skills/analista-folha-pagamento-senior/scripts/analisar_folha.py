#!/usr/bin/env python3
"""
analisar_folha.py — Motor de cálculo da skill analista-folha-pagamento-senior.

Lê a planilha consolidada de holerites (abas "Resumo" e "Rubricas", no formato
gerado pela skill consolidar-holerites-aw) e gera um Excel de auditoria com:
  - Resumo Executivo
  - Rubricas Atípicas (desvio estatístico por função + rubrica)
  - Média Salarial por Função
  - Custo Empresa (FGTS + CPRB/CPP híbrido de desoneração + RAT/FAP + Terceiros)
  - Premissas (memória de cálculo — leia antes de comentar qualquer número)

Não inventa dado: coluna ausente ou insuficiente gera aviso explícito na aba
Premissas, nunca um número chutado. Ver references/02-encargos-desoneracao.md
para a fonte de cada alíquota.
"""
import argparse
import re
import sys
import warnings
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

warnings.filterwarnings("ignore")

# ---------------------------------------------------------------------------
# Constantes de encargos — ver references/02-encargos-desoneracao.md
# ---------------------------------------------------------------------------

CPRB_ALIQUOTA_CHEIA = 0.045  # CPRB cheia s/ receita bruta, construção civil (Lei 12.546/2011)
FGTS_PCT = 0.08

# RAT: a faixa legal é 1% a 3% por grau de risco do CNAE. Construção civil
# (CNAE 41 a 43) é grau de risco 3 -> 3%. NÃO use ponto médio da faixa: para
# esta empresa, 2% subestima o custo patronal em 1 p.p. da base do INSS e
# contamina toda decisão de preço. O valor real vem do CNAE do estabelecimento
# e é sobrescrito por --rat.
RAT_CONSTRUCAO_CIVIL = 0.03

# Terceiros/Sistema S: varia pelo código FPAS do estabelecimento. 5,8% é o
# valor usual do FPAS 507 (construção civil) e é apenas o ponto de partida —
# sobrescreva com --terceiros quando a guia da empresa mostrar outro.
TERCEIROS_PADRAO = 0.058

FAP_MIN, FAP_MAX = 0.5, 2.0

# Cronograma de reoneração gradual (Lei 14.973/2024): percentual da CPRB cheia
# ainda aplicável sobre a receita bruta, e CPP básica sobre a folha, por ano.
# 2025 tem fontes secundárias divergentes sobre a CPP (5% x 10%) — sinalizado
# no aviso; confirme na lei antes de fechar número para esse ano.
CRONOGRAMA_DESONERACAO = {
    2025: {"pct_cprb_aplicavel": 0.80, "cpp_pct_folha": 0.05,
           "aviso": "Fontes secundárias divergem sobre a CPP de 2025 (5% vs 10%). "
                    "Confirme na Lei 14.973/2024 antes de fechar o número para esse ano."},
    2026: {"pct_cprb_aplicavel": 0.60, "cpp_pct_folha": 0.10, "aviso": None},
    2027: {"pct_cprb_aplicavel": 0.40, "cpp_pct_folha": 0.15, "aviso": None},
    2028: {"pct_cprb_aplicavel": 0.00, "cpp_pct_folha": 0.20,
           "aviso": "Cronograma encerrado: CPRB extinta, CPP volta a 20% cheio "
                    "sobre a folha (INSS patronal integral)."},
}

RESUMO_REQUIRED = [
    "Competência", "Cargo", "Salário Base", "Sal. Contr. INSS",
    "Total Vencimentos", "Base Cálc. FGTS", "FGTS do Mês",
    "Nome do Funcionário", "Código Funcionário",
]
RUBRICAS_REQUIRED = [
    "Competência", "Código Funcionário", "Nome do Funcionário",
    "Descrição Rubrica", "Natureza", "Vencimento", "Desconto",
]

MESES_PT = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
            "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12}


# ---------------------------------------------------------------------------
# Carga e normalização
# ---------------------------------------------------------------------------

def parse_competencia(valor):
    """Extrai (ano, mes, chave_ordenavel) de um valor de Competência.
    Aceita 'MM/AAAA', 'AAAA-MM', datas do pandas/Excel, ou nome de mês em
    PT-BR ('Janeiro/2026'). Não interpretou -> (None, None, str(valor)), e
    quem chamar trata como não-cronológico em vez de arriscar ordenação errada."""
    if pd.isna(valor):
        return None, None, ""
    if isinstance(valor, (pd.Timestamp, datetime)):
        return valor.year, valor.month, f"{valor.year:04d}-{valor.month:02d}"

    s = str(valor).strip()

    m = re.match(r"^(\d{1,2})[/\-](\d{4})$", s)
    if m:
        mes, ano = int(m.group(1)), int(m.group(2))
        if 1 <= mes <= 12:
            return ano, mes, f"{ano:04d}-{mes:02d}"

    m = re.match(r"^(\d{4})[/\-](\d{1,2})$", s)
    if m:
        ano, mes = int(m.group(1)), int(m.group(2))
        if 1 <= mes <= 12:
            return ano, mes, f"{ano:04d}-{mes:02d}"

    m = re.match(r"^([A-Za-zçÇ]+)[/\- ](\d{4})$", s)
    if m:
        nome, ano = m.group(1).lower()[:3], int(m.group(2))
        if nome in MESES_PT:
            return ano, MESES_PT[nome], f"{ano:04d}-{MESES_PT[nome]:02d}"

    return None, None, s


def carregar_base(caminho):
    try:
        resumo = pd.read_excel(caminho, sheet_name="Resumo")
    except Exception as e:
        sys.exit(f"[erro] Não consegui ler a aba 'Resumo' de {caminho}: {e}")
    try:
        rubricas = pd.read_excel(caminho, sheet_name="Rubricas")
    except Exception as e:
        sys.exit(f"[erro] Não consegui ler a aba 'Rubricas' de {caminho}: {e}")

    faltando_resumo = [c for c in RESUMO_REQUIRED if c not in resumo.columns]
    faltando_rubricas = [c for c in RUBRICAS_REQUIRED if c not in rubricas.columns]
    if faltando_resumo:
        sys.exit(
            f"[erro] Aba 'Resumo' sem as colunas esperadas: {faltando_resumo}. "
            "Confirme se o arquivo veio da skill consolidar-holerites-aw, ou ajuste "
            "os nomes das colunas antes de rodar — não estime o que falta."
        )
    if faltando_rubricas:
        sys.exit(f"[erro] Aba 'Rubricas' sem as colunas esperadas: {faltando_rubricas}.")

    resumo["Código Funcionário"] = resumo["Código Funcionário"].astype(str).str.strip()
    rubricas["Código Funcionário"] = rubricas["Código Funcionário"].astype(str).str.strip()
    resumo["Cargo"] = resumo["Cargo"].fillna("(Cargo não identificado)").astype(str).str.strip()

    for df in (resumo, rubricas):
        anos, meses, chaves = [], [], []
        for v in df["Competência"]:
            a, m, k = parse_competencia(v)
            anos.append(a)
            meses.append(m)
            chaves.append(k)
        df["_ano"], df["_mes"], df["_competencia_chave"] = anos, meses, chaves

    return resumo, rubricas


# ---------------------------------------------------------------------------
# 1) Rubricas atípicas
# ---------------------------------------------------------------------------

def calcular_rubricas_atipicas(resumo, rubricas):
    chave_resumo = resumo[["Código Funcionário", "_competencia_chave", "Cargo"]].drop_duplicates()
    rub = rubricas.merge(chave_resumo, on=["Código Funcionário", "_competencia_chave"], how="left")
    rub["Cargo"] = rub["Cargo"].fillna("(Cargo não identificado)")

    rub["Natureza_norm"] = rub["Natureza"].astype(str).str.strip().str.lower()
    proventos = rub[rub["Natureza_norm"].str.startswith("prov", na=False)].copy()
    proventos["Valor"] = pd.to_numeric(proventos["Vencimento"], errors="coerce").fillna(0.0)

    agrupado = (proventos.groupby(["Cargo", "Descrição Rubrica", "_competencia_chave"])["Valor"]
                .sum().reset_index())

    linhas = []
    for (cargo, rubrica_desc), grupo in agrupado.groupby(["Cargo", "Descrição Rubrica"]):
        grupo = grupo.sort_values("_competencia_chave")
        valores = grupo["Valor"].tolist()
        competencias = grupo["_competencia_chave"].tolist()
        if not valores:
            continue

        ultimo_valor, ultima_competencia, n_meses = valores[-1], competencias[-1], len(valores)
        atipico, motivo = False, ""

        if n_meses >= 4:
            historico = valores[:-1]
            media_hist = float(np.mean(historico))
            desvio_hist = float(np.std(historico, ddof=1)) if len(historico) >= 2 else 0.0
            if desvio_hist > 0:
                z = (ultimo_valor - media_hist) / desvio_hist
                if abs(z) > 2:
                    atipico = True
                    motivo = f"{z:+.1f} desvios-padrão da média histórica (R$ {media_hist:,.2f})"
        if not atipico and n_meses >= 2:
            anterior = valores[-2]
            if anterior != 0:
                variacao = (ultimo_valor - anterior) / abs(anterior)
                if abs(variacao) > 0.30:
                    atipico = True
                    motivo = (f"variação de {variacao:+.0%} sobre o mês anterior "
                              f"(R$ {anterior:,.2f}); histórico curto ({n_meses} mês(es))")

        linhas.append({
            "Cargo": cargo,
            "Descrição Rubrica": rubrica_desc,
            "Última Competência": ultima_competencia,
            "Valor no Mês (R$)": round(ultimo_valor, 2),
            "Meses no Histórico": n_meses,
            "Atípico": "Sim" if atipico else "Não",
            "Motivo": motivo,
        })

    df = pd.DataFrame(linhas)
    if not df.empty:
        df = df.sort_values(["Atípico", "Valor no Mês (R$)"], ascending=[False, False])
    return df


# ---------------------------------------------------------------------------
# 2) Média salarial por função
# ---------------------------------------------------------------------------

def calcular_media_salarial(resumo):
    ultima_competencia = resumo["_competencia_chave"].max()
    base = resumo[resumo["_competencia_chave"] == ultima_competencia].copy()
    if base.empty:
        base = resumo.copy()

    for col in ["Salário Base", "Total Vencimentos"]:
        base[col] = pd.to_numeric(base[col], errors="coerce")

    agg = base.groupby("Cargo").agg(
        Funcionários=("Código Funcionário", "nunique"),
        Salário_Base_Médio=("Salário Base", "mean"),
        Salário_Base_Mediana=("Salário Base", "median"),
        Salário_Base_DesvPad=("Salário Base", "std"),
        Total_Vencimentos_Médio=("Total Vencimentos", "mean"),
    ).reset_index().rename(columns={
        "Salário_Base_Médio": "Salário Base — Média (R$)",
        "Salário_Base_Mediana": "Salário Base — Mediana (R$)",
        "Salário_Base_DesvPad": "Salário Base — Desvio-Padrão (R$)",
        "Total_Vencimentos_Médio": "Total Vencimentos — Média (R$)",
    })
    for c in agg.columns:
        if c not in ("Cargo", "Funcionários"):
            agg[c] = agg[c].round(2)
    agg = agg.sort_values("Total Vencimentos — Média (R$)", ascending=False)
    agg["Competência de Referência"] = ultima_competencia
    return agg


# ---------------------------------------------------------------------------
# 3) Custo empresa
# ---------------------------------------------------------------------------

def calcular_custo_empresa(resumo, receita_bruta, fap, ano_cprb, rat, rat_informado,
                           terceiros, terceiros_informado):
    ultima_competencia = resumo["_competencia_chave"].max()
    base = resumo[resumo["_competencia_chave"] == ultima_competencia].copy()
    if base.empty:
        base = resumo.copy()

    for col in ["Total Vencimentos", "Sal. Contr. INSS", "Base Cálc. FGTS", "FGTS do Mês"]:
        base[col] = pd.to_numeric(base[col], errors="coerce").fillna(0.0)

    n_func = base["Código Funcionário"].nunique()
    folha_bruta = base["Total Vencimentos"].sum()
    base_inss = base["Sal. Contr. INSS"].sum()
    base_fgts = base["Base Cálc. FGTS"].sum()
    fgts_apurado = base["FGTS do Mês"].sum()

    schedule = CRONOGRAMA_DESONERACAO.get(ano_cprb)
    if schedule is None:
        aviso_schedule = (
            f"Ano {ano_cprb} fora do cronograma de reoneração conhecido (2025-2028). "
            "Confirme a alíquota vigente — nenhuma CPP/CPRB foi calculada."
        )
        cpp_pct_folha, pct_cprb_aplicavel = 0.0, 0.0
    else:
        cpp_pct_folha = schedule["cpp_pct_folha"]
        pct_cprb_aplicavel = schedule["pct_cprb_aplicavel"]
        aviso_schedule = schedule["aviso"]

    cpp = base_inss * cpp_pct_folha
    rat_valor = base_inss * rat * fap
    terceiros_valor = base_inss * terceiros

    cprb_reais = receita_bruta * CPRB_ALIQUOTA_CHEIA * pct_cprb_aplicavel \
        if receita_bruta is not None else None

    fgts_usado = fgts_apurado if fgts_apurado > 0 else base_fgts * FGTS_PCT
    custo_sobre_folha = fgts_usado + cpp + rat_valor + terceiros_valor
    custo_total_mes = folha_bruta + custo_sobre_folha
    custo_medio_por_cabeca = (custo_total_mes / n_func) if n_func else None

    resumo_empresa = {
        "Competência de Referência": ultima_competencia,
        "Nº Funcionários": n_func,
        "Folha Bruta (R$)": round(folha_bruta, 2),
        "FGTS (R$)": round(fgts_usado, 2),
        f"CPP {cpp_pct_folha:.0%} sobre a folha (R$)": round(cpp, 2),
        f"RAT {rat:.1%} x FAP {fap} (R$)": round(rat_valor, 2),
        f"Terceiros / Sistema S {terceiros:.1%} (R$)": round(terceiros_valor, 2),
        "Custo Total Folha + Encargos (R$)": round(custo_total_mes, 2),
        "Fator sobre a Folha Bruta": (round(custo_total_mes / folha_bruta, 4)
                                     if folha_bruta else None),
        "Custo Médio por Funcionário (R$)": (round(custo_medio_por_cabeca, 2)
                                            if custo_medio_por_cabeca else None),
        "CPRB sobre Receita Bruta (R$)": (round(cprb_reais, 2) if cprb_reais is not None
                                          else "N/D — informe --receita-bruta"),
        "Alíquota CPRB aplicável no ano": (
            f"{CPRB_ALIQUOTA_CHEIA * pct_cprb_aplicavel:.2%} da receita bruta"
            if schedule else "N/D"),
    }

    por_funcao = base.groupby("Cargo").agg(
        Funcionários=("Código Funcionário", "nunique"),
        Salário_Bruto=("Total Vencimentos", "sum"),
    ).reset_index()
    if folha_bruta > 0:
        por_funcao["Encargos Rateados (R$)"] = (
            por_funcao["Salário_Bruto"] / folha_bruta) * custo_sobre_folha
    else:
        por_funcao["Encargos Rateados (R$)"] = 0.0
    por_funcao["Custo Total da Função (R$)"] = (
        por_funcao["Salário_Bruto"] + por_funcao["Encargos Rateados (R$)"])
    por_funcao["Custo Médio por Cabeça (R$)"] = (
        por_funcao["Custo Total da Função (R$)"] / por_funcao["Funcionários"])
    por_funcao = por_funcao.rename(columns={"Salário_Bruto": "Salário Bruto (R$)"})
    for c in ["Salário Bruto (R$)", "Encargos Rateados (R$)",
              "Custo Total da Função (R$)", "Custo Médio por Cabeça (R$)"]:
        por_funcao[c] = por_funcao[c].round(2)
    por_funcao = por_funcao.sort_values("Custo Total da Função (R$)", ascending=False)

    rotulo_rat = "confirmado (--rat)" if rat_informado else (
        "PREMISSA: grau de risco 3 da construção civil (CNAE 41 a 43). "
        "Confirme o RAT do CNAE do estabelecimento")
    rotulo_terceiros = "confirmado (--terceiros)" if terceiros_informado else (
        "PREMISSA: usual do FPAS 507. Confirme na guia da empresa")

    premissas = [
        f"Competência de referência: {ultima_competencia}",
        f"Ano usado para o cronograma de desoneração: {ano_cprb}",
        (f"CPP sobre a folha: {cpp_pct_folha:.0%} | CPRB aplicável: {pct_cprb_aplicavel:.0%} da "
         f"alíquota cheia de {CPRB_ALIQUOTA_CHEIA:.1%} sobre a receita bruta "
         f"(construção civil, Lei 12.546/2011)"),
        f"RAT usado: {rat:.2%} — {rotulo_rat}",
        f"FAP usado: {fap} — {'confirmado (--fap)' if fap != 1.0 else 'PREMISSA: 1,0 (neutro). Informe o FAP publicado da empresa'}",
        f"Terceiros/Sistema S: {terceiros:.2%} — {rotulo_terceiros}",
        ("FGTS: usado o valor já apurado no holerite (coluna 'FGTS do Mês') quando disponível; "
         "senão, 8% sobre a Base Cálc. FGTS"),
        ("Provisão de multa rescisória (40% do FGTS + aviso prévio + férias/13º proporcionais) "
         "NÃO foi calculada: depende da rotatividade real. Peça a taxa de rotatividade antes de "
         "estimar; não use percentual genérico."),
        (f"Receita bruta informada: "
         f"{'Sim — R$ ' + format(receita_bruta, ',.2f') if receita_bruta is not None else 'Não — CPRB em R$ não calculada, apenas o percentual aplicável'}"),
    ]
    if aviso_schedule:
        premissas.append(f"AVISO: {aviso_schedule}")
    if not rat_informado or fap == 1.0 or not terceiros_informado:
        premissas.append(
            "ATENÇÃO: o custo empresa acima contém premissa de encargo não confirmada. "
            "Não use este número para decisão de preço de venda, orçamento ou corte de "
            "quadro antes de confirmar RAT, FAP e Terceiros com a guia e o FAP publicado."
        )

    return resumo_empresa, por_funcao, premissas


# ---------------------------------------------------------------------------
# Escrita do Excel de saída
# ---------------------------------------------------------------------------

def autofit(ws, max_width=60):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[letter].width = min(max(length + 2, 10), max_width)


def escrever_df(wb, nome_aba, df, nome_tabela):
    ws = wb.create_sheet(nome_aba)
    if df is None or df.empty:
        ws.append(["(sem dados suficientes para esta aba)"])
        return ws
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for _, row in df.iterrows():
        ws.append(list(row))
    ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
    tabela = Table(displayName=nome_tabela, ref=ref)
    tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabela)
    ws.freeze_panes = "A2"
    autofit(ws)
    return ws


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--input", required=True, help="Excel consolidado (abas Resumo e Rubricas).")
    ap.add_argument("--output", required=True, help="Excel de auditoria a ser gerado.")
    ap.add_argument("--receita-bruta", type=float, default=None,
                    help="Receita bruta do período, para calcular a CPRB em R$ (opcional).")
    ap.add_argument("--fap", type=float, default=1.0,
                    help=f"Fator Acidentário de Prevenção, entre {FAP_MIN} e {FAP_MAX} "
                         "(padrão 1.0 = premissa neutra).")
    ap.add_argument("--rat", type=float, default=None,
                    help="RAT em decimal (0.01, 0.02 ou 0.03). Padrão: 0.03, grau de risco 3 "
                         "da construção civil (CNAE 41 a 43).")
    ap.add_argument("--terceiros", type=float, default=None,
                    help="Terceiros/Sistema S em decimal. Padrão: 0.058 (usual do FPAS 507).")
    ap.add_argument("--ano-cprb", type=int, default=None,
                    help="Ano do cronograma de desoneração (padrão: ano da última competência).")
    args = ap.parse_args()

    if not FAP_MIN <= args.fap <= FAP_MAX:
        sys.exit(f"[erro] FAP {args.fap} fora da faixa legal {FAP_MIN} a {FAP_MAX}.")
    if args.rat is not None and not 0.01 <= args.rat <= 0.03:
        sys.exit(f"[erro] RAT {args.rat} fora da faixa legal 0.01 a 0.03.")

    rat_informado = args.rat is not None
    rat = args.rat if rat_informado else RAT_CONSTRUCAO_CIVIL
    terceiros_informado = args.terceiros is not None
    terceiros = args.terceiros if terceiros_informado else TERCEIROS_PADRAO

    resumo, rubricas = carregar_base(args.input)

    anos_validos = [a for a in resumo["_ano"].tolist() if a is not None]
    ano_cprb = args.ano_cprb or (max(anos_validos) if anos_validos else datetime.now().year)

    print(f"[info] {resumo['Código Funcionário'].nunique()} funcionário(s), "
          f"{resumo['_competencia_chave'].nunique()} competência(s) na base.")

    df_atipicas = calcular_rubricas_atipicas(resumo, rubricas)
    df_media_salarial = calcular_media_salarial(resumo)
    resumo_empresa, df_por_funcao, premissas = calcular_custo_empresa(
        resumo, args.receita_bruta, args.fap, ano_cprb,
        rat, rat_informado, terceiros, terceiros_informado)

    n_atipicas = int((df_atipicas["Atípico"] == "Sim").sum()) if not df_atipicas.empty else 0
    print(f"[info] {n_atipicas} combinação(ões) função+rubrica sinalizada(s) como atípica.")
    print(f"[info] Custo total (folha + encargos): "
          f"R$ {resumo_empresa['Custo Total Folha + Encargos (R$)']:,.2f}")
    if not rat_informado:
        print("[aviso] RAT não informado: usando 3% (grau de risco 3, construção civil). "
              "Confirme o CNAE do estabelecimento.")

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Resumo Executivo")
    ws.append(["Auditoria de Folha de Pagamento — Resumo Executivo"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
    for k, v in resumo_empresa.items():
        ws.append([k, v])
    ws.append([])
    ws.append([f"Top rubricas atípicas (competência {resumo_empresa['Competência de Referência']})"])
    ws[ws.max_row][0].font = Font(bold=True)
    ws.append(["Cargo", "Rubrica", "Valor no Mês (R$)", "Motivo"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    if not df_atipicas.empty:
        for _, row in df_atipicas[df_atipicas["Atípico"] == "Sim"].head(10).iterrows():
            ws.append([row["Cargo"], row["Descrição Rubrica"],
                       row["Valor no Mês (R$)"], row["Motivo"]])
    autofit(ws, max_width=80)

    escrever_df(wb, "Rubricas Atípicas", df_atipicas, "tbAtipicas")
    escrever_df(wb, "Média Salarial por Função", df_media_salarial, "tbMediaSalarial")

    ws = wb.create_sheet("Custo Empresa")
    ws.append(["Custo Empresa Consolidado"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for k, v in resumo_empresa.items():
        ws.append([k, v])
    ws.append([])
    ws.append(["Custo por Função"])
    ws[ws.max_row][0].font = Font(bold=True)
    header_row = ws.max_row + 1
    for col_i, col_name in enumerate(df_por_funcao.columns, start=1):
        ws.cell(row=header_row, column=col_i, value=col_name).font = Font(bold=True)
    for _, row in df_por_funcao.iterrows():
        ws.append(list(row))
    autofit(ws, max_width=60)

    ws = wb.create_sheet("Premissas")
    ws.append(["Premissas e Memória de Cálculo — leia antes de comentar qualquer número"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    for p in premissas:
        ws.append([p])
        ws[ws.max_row][0].alignment = Alignment(wrap_text=True, vertical="top")
    autofit(ws, max_width=110)

    wb.save(args.output)
    print(f"[info] Excel de auditoria salvo em: {args.output}")


if __name__ == "__main__":
    main()
