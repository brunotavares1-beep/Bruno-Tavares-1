#!/usr/bin/env python3
"""
Extrai o relatório "RESUMO DA FOLHA" (TOTVS RM) de 1..N PDFs e monta um Excel
com 6 abas: Resumo (layout aprovado pelo usuário, custo total da folha por
competência), Consolidado, Proventos, Descontos, Informativa e
Encargos INSS-FGTS-IRRF (inclui IRRF e Situações/headcount).

Uso (unitário ou lote — PDFs, pastas e ZIPs podem ser misturados):
    python3 extrair_resumo_folha.py <pdf|pasta|zip> [...] --out <saida.xlsx>
                                    [--categorias categorias_rubricas.csv]

Nunca inventa valor: toda linha útil do PDF vira dado na planilha ou
pendência (aba Consolidado). Nada é descartado em silêncio.
"""
import argparse
import csv
import re
import sys
import tempfile
import zipfile
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Cabeçalho/rodapé repetido de cada página. ATENÇÃO: "Empresa:" só é cabeçalho
# quando seguido de "<código> - <razão social>". No bloco de encargos existe a
# linha "Empresa: 76.553,35 Valor FGTS Rescisório: ..." (INSS patronal) que
# NÃO pode ser descartada — por isso regex estrito em vez de startswith.
HEADER_RES = [re.compile(p) for p in (
    r"^Empresa:\s*\d+\s*-\s*\D", r"^CNPJ:", r"^Cálculo:\s*\D", r"^Competência:\s*\d{2}/\d{4}",
    r"^Complemento de cálculo:", r"^RESUMO DA FOLHA$", r"^Rubrica Nome da Rubrica",
    r"^Sistema licenciado", r"^Página:\s*\d+/\d+",
)]


def is_header(line):
    return any(r.search(line) for r in HEADER_RES)
SECTION_MARKERS = {"Folha Mensal", "PROVENTOS", "DESCONTOS", "INFORMATIVA"}
MONEY = r"-?\d{1,3}(?:\.\d{3})*,\d{2}"
VALUE_ANY = rf"({MONEY}|\d+)"
PAIR_RE = re.compile(rf"([^:]+):\s*{VALUE_ANY}")


def to_float(s):
    s = s.strip().rstrip("*")
    return float(s.replace(".", "").replace(",", "."))


def is_money_token(t):
    return bool(re.fullmatch(rf"{MONEY}\*?", t))


def is_int_token(t):
    return bool(re.fullmatch(r"\d+", t))


def extrair_meta(primeira_pagina_texto):
    meta = {}
    for line in primeira_pagina_texto.split("\n"):
        line = line.strip()
        if line.startswith("Empresa:"):
            meta["empresa"] = line.split("Empresa:", 1)[1].split("Página:")[0].strip()
        elif line.startswith("CNPJ:"):
            resto = line.split("CNPJ:", 1)[1]
            partes = resto.split("Emissão:")
            meta["cnpj"] = partes[0].strip()
            if len(partes) > 1:
                meta["emissao"] = partes[1].strip()
        elif line.startswith("Cálculo:"):
            resto = line.split("Cálculo:", 1)[1]
            partes = resto.split("Hora:")
            meta["calculo"] = partes[0].strip()
        elif line.startswith("Competência:"):
            meta["competencia"] = line.split("Competência:", 1)[1].strip()
        elif line.startswith("Complemento de cálculo:"):
            meta["complemento"] = line.split("Complemento de cálculo:", 1)[1].strip()
    return meta


def parse_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        paginas_texto = [p.extract_text() or "" for p in pdf.pages]

    meta = extrair_meta(paginas_texto[0]) if paginas_texto else {}

    all_lines = []
    for texto in paginas_texto:
        for line in texto.split("\n"):
            line = line.strip()
            if not line or is_header(line):
                continue
            all_lines.append(line)

    section = None
    rubricas = {"PROVENTOS": [], "DESCONTOS": [], "INFORMATIVA": []}
    totais_impressos = {}
    liquido_geral = None
    linhas_pos_folha = []  # tudo depois do fim das 3 seções (encargos/IRRF/situações)
    fim_folha = False

    for line in all_lines:
        if not fim_folha and line in SECTION_MARKERS:
            if line != "Folha Mensal":
                section = line
            continue

        m_total = re.fullmatch(rf"Total:\s*({MONEY})", line)
        if not fim_folha and m_total:
            totais_impressos[section] = to_float(m_total.group(1))
            continue

        m_liq = re.fullmatch(rf"Líquido Geral:\s*({MONEY})", line)
        if m_liq:
            liquido_geral = to_float(m_liq.group(1))
            fim_folha = True
            continue

        if line == "Total: Resumo Geral Mensal e Complementar":
            fim_folha = True
            continue

        if fim_folha:
            linhas_pos_folha.append(line)
            continue

        tokens = line.split(" ")
        if (
            section in rubricas
            and len(tokens) >= 4
            and tokens[0].isdigit()
            and is_money_token(tokens[-1])
            and is_money_token(tokens[-2])
            and is_int_token(tokens[-3])
        ):
            codigo = tokens[0]
            nome = " ".join(tokens[1:-3]).strip()
            nemp = int(tokens[-3])
            vinf = to_float(tokens[-2])
            vcalc_raw = tokens[-1]
            informativa_marker = vcalc_raw.endswith("*")
            vcalc = to_float(vcalc_raw)
            rubricas[section].append(
                {
                    "codigo": codigo,
                    "nome": nome,
                    "n_empregados": nemp,
                    "valor_informado": vinf,
                    "valor_calculado": vcalc,
                    "informativa": informativa_marker or section == "INFORMATIVA",
                }
            )
        else:
            rubricas.setdefault("_PENDENCIAS", []).append(f"[{section}] {line}")

    # --- Encargos / IRRF / Situações: pares "Label: valor" (2 por linha) ---
    encargos = {}
    irrf_calculo = {}
    irrf_pagamento = {}
    situacoes = {}

    bloco = None  # None -> encargos ; "irrf" ; "situacoes"
    for line in linhas_pos_folha:
        if line.startswith("IRRF conforme competência"):
            bloco = "irrf"
            continue
        if line == "Situações":
            bloco = "situacoes"
            continue
        if line == "INSS FGTS, PIS e ISS":
            continue

        pares = PAIR_RE.findall(line)
        sobra = PAIR_RE.sub("", line).strip()
        if not pares or sobra:
            # linha não lida por inteiro -> pendência (nunca descarta texto)
            rubricas.setdefault("_PENDENCIAS", []).append(f"[RESUMO] {line}")
            if not pares:
                continue

        for i, (label, valor) in enumerate(pares):
            label = label.strip()
            if bloco == "situacoes":
                situacoes[label] = int(valor)
            elif bloco == "irrf":
                destino = irrf_calculo if i == 0 else irrf_pagamento
                destino[label] = to_float(valor)
            else:
                encargos[label] = to_float(valor)

    return {
        "arquivo": Path(pdf_path).name,
        "n_linhas_uteis": len(all_lines),
        "meta": meta,
        "rubricas": rubricas,
        "totais_impressos": totais_impressos,
        "liquido_geral": liquido_geral,
        "encargos": encargos,
        "irrf_calculo": irrf_calculo,
        "irrf_pagamento": irrf_pagamento,
        "situacoes": situacoes,
    }


def carregar_categorias(csv_path):
    mapa = {}
    if csv_path and Path(csv_path).exists():
        with open(csv_path, encoding="utf-8") as f:
            for row in csv.DictReader(f):
                mapa[row["codigo"].strip()] = row["categoria"].strip()
    return mapa


# ------------------------------------------------------------ Entrada em lote
def coletar_pdfs(entradas, tmpdir):
    """Aceita PDFs, pastas (recursivo) e ZIPs. Retorna lista de caminhos de PDF."""
    pdfs = []
    for e in entradas:
        p = Path(e)
        if p.is_dir():
            pdfs += sorted(x for x in p.rglob("*") if x.suffix.lower() == ".pdf")
        elif p.suffix.lower() == ".zip":
            destino = Path(tmpdir) / p.stem
            with zipfile.ZipFile(p) as z:
                z.extractall(destino)
            pdfs += coletar_pdfs([destino], tmpdir)
        elif p.suffix.lower() == ".pdf":
            pdfs.append(p)
    return pdfs


def chave_comp(comp):
    m = re.fullmatch(r"(\d{2})/(\d{4})", comp or "")
    return (int(m.group(2)), int(m.group(1))) if m else (9999, 99)


# ------------------------------------------------------------------ Layout ---
# Linhas fixas da 1ª aba (layout aprovado pelo usuário - "Exemplo - Extração").
CATS_PROVENTOS = ["Salário Base/Dias", "Horas Extras", "Adicional Noturno", "Produção/Tarefa",
                  "Prêmio/Bônus", "Férias", "13º Salário", "Rescisão", "Provisão/Estorno",
                  "Benefícios", "Sindicato", "Diferença/Ajuste", "Outros/Diversos"]
CATS_DESCONTOS = ["INSS/IRRF Empregado", "Adiantamento Salarial", "Empréstimo Consignado",
                  "Benefícios", "Retenção Judicial", "Sindicato", "Rescisão", "Férias",
                  "13º Salário", "Salário Base/Dias", "Provisão/Estorno", "Prêmio/Bônus",
                  "Diferença/Ajuste", "Custo Empresa (referência)", "Outros/Diversos"]
CATS_INFORMATIVA = ["FGTS (Informativa)", "Rescisão", "Férias", "13º Salário",
                    "Benefícios (Informativa)"]
# (rótulo na aba Resumo, label exato no PDF / aba Encargos)
ENCARGOS_RESUMO = [
    ("FGTS (mensal)", "Valor do FGTS"),
    ("FGTS Aprendiz", "Valor do FGTS Aprendiz"),
    ("FGTS Rescisório (8% + multa 40%)", "Valor FGTS Rescisório"),
    ("FGTS Rescisório mês anterior", "Valor FGTS Resc. mês ant."),
    ("INSS Patronal – empregados (reoneração)", "Empresa"),
    ("INSS Patronal – contribuintes/pró-labore", "Contribuintes"),
    ("RAT / FAP", "RAT"),
    ("Terceiros (Sistema S)", "Terceiros"),
]

SH_RES, SH_CON, SH_PRO, SH_DES, SH_INF, SH_ENC = (
    "Resumo", "Consolidado", "Proventos", "Descontos", "Informativa", "Encargos INSS-FGTS-IRRF")

FONT = "Arial"
F_TITLE = Font(name=FONT, bold=True, size=13, color="1F4E78")
F_HDR = Font(name=FONT, bold=True, color="FFFFFF")
F_SEC = Font(name=FONT, bold=True, color="1F4E78")
F_BOLD = Font(name=FONT, bold=True)
F_NORM = Font(name=FONT)
F_NOTE = Font(name=FONT, italic=True, size=9, color="595959")
FILL_HDR = PatternFill("solid", fgColor="1F4E78")
FILL_SEC = PatternFill("solid", fgColor="DDEBF7")
FILL_TOT = PatternFill("solid", fgColor="F2F2F2")
FILL_KEY = PatternFill("solid", fgColor="FFF2CC")
FILL_ALERT = PatternFill("solid", fgColor="FFC7CE")
FILL_NEW = PatternFill("solid", fgColor="E2EFDA")
NUM = '#,##0.00;[Red]-#,##0.00;"-"'
INT = '#,##0;[Red]-#,##0;"-"'
THIN = Side(style="thin", color="BFBFBF")


def _hdr(ws, row, values, col0=1):
    for j, v in enumerate(values):
        c = ws.cell(row=row, column=col0 + j, value=v)
        c.font, c.fill = F_HDR, FILL_HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _q(sh):
    return f"'{sh}'" if " " in sh or "-" in sh else sh


# ------------------------------------------------------------- Abas base ---
def aba_rubricas(wb, nome, lista, bloco, categorias):
    ws = wb.create_sheet(nome)
    _hdr(ws, 1, ["Competência", "Código", "Nome da Rubrica", "Categoria",
                 "Nº Empregados/Contribuintes", "Valor Informado", "Valor Calculado",
                 "Informativa (*) — não compõe total", "Arquivo"])
    r = 2
    for d in lista:
        comp = d["meta"]["competencia"]
        for it in d["rubricas"].get(bloco, []):
            cat = categorias.get(it["codigo"], "NÃO CLASSIFICADO")
            vals = [comp, it["codigo"], it["nome"], cat, it["n_empregados"],
                    it["valor_informado"], it["valor_calculado"],
                    "SIM" if it["informativa"] else "", d["arquivo"]]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.font = F_NORM
                if j in (6, 7):
                    c.number_format = NUM
                if cat == "NÃO CLASSIFICADO":
                    c.fill = FILL_ALERT
            r += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(r - 1, 1)}"
    for col, w in zip("ABCDEFGHI", (12, 8, 46, 26, 14, 16, 16, 18, 30)):
        ws.column_dimensions[col].width = w
    return ws


def _union(lista, chave):
    labels = []
    for d in lista:
        for k in d[chave]:
            if k not in labels:
                labels.append(k)
    return labels


def aba_encargos(wb, lista):
    """Formato largo: A = indicador, B.. = uma coluna por competência (mesma
    ordem das colunas da aba Resumo). Retorna {label_encargo: linha}."""
    ws = wb.create_sheet(SH_ENC)
    comps = [d["meta"]["competencia"] for d in lista]
    _hdr(ws, 1, ["Indicador"] + comps)
    linhas = {}
    r = 2

    def secao(titulo):
        nonlocal r
        c = ws.cell(row=r, column=1, value=titulo)
        c.font, c.fill = F_SEC, FILL_SEC
        for j in range(2, len(comps) + 2):
            ws.cell(row=r, column=j).fill = FILL_SEC
        r += 1

    def bloco(chave, prefixo, fmt):
        nonlocal r
        for lab in _union(lista, chave):
            ws.cell(row=r, column=1, value=f"{prefixo}{lab}").font = F_NORM
            for j, d in enumerate(lista, 2):
                c = ws.cell(row=r, column=j, value=d[chave].get(lab))
                c.number_format, c.font = fmt, F_NORM
            if chave == "encargos":
                linhas[lab] = r
            r += 1

    secao("INSS / FGTS / PIS / ISS")
    bloco("encargos", "", NUM)
    secao("IRRF — conforme competência do CÁLCULO")
    bloco("irrf_calculo", "", NUM)
    secao("IRRF — conforme competência do PAGAMENTO")
    bloco("irrf_pagamento", "", NUM)
    secao("SITUAÇÕES (headcount)")
    bloco("situacoes", "", INT)
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 40
    for j in range(2, len(comps) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 16
    return ws, linhas


def aba_consolidado(wb, lista, categorias):
    """A=Competência B..I = campos usados pela aba Resumo (não mover)."""
    ws = wb.create_sheet(SH_CON)
    heads = ["Competência", "Total Proventos (PDF)", "Total Descontos (PDF)",
             "Total Informativa (PDF)", "Líquido Geral (PDF)", "Custo Total Empresa",
             "Nº Trabalhando", "Admissões", "Demitido", "Nº Empregados", "Empresa", "CNPJ",
             "Cálculo", "Emissão", "Arquivo", "Linhas Proventos", "Linhas Descontos",
             "Linhas Informativa", "Rubricas NÃO CLASSIFICADAS", "Pendências", "Status"]
    _hdr(ws, 1, heads)
    for i, d in enumerate(lista, 2):
        s, t, m, rb = d["situacoes"], d["totais_impressos"], d["meta"], d["rubricas"]
        nc = sum(1 for b in ("PROVENTOS", "DESCONTOS", "INFORMATIVA")
                 for it in rb.get(b, []) if it["codigo"] not in categorias)
        pend = len(rb.get("_PENDENCIAS", []))
        vals = [m["competencia"], t.get("PROVENTOS"), t.get("DESCONTOS"), t.get("INFORMATIVA"),
                d["liquido_geral"], None, s.get("Trabalhando"), s.get("Admissões"),
                s.get("Demitido"), s.get("No. Empregados"), m.get("empresa"), m.get("cnpj"),
                m.get("calculo"), m.get("emissao"), d["arquivo"],
                len(rb["PROVENTOS"]), len(rb["DESCONTOS"]), len(rb["INFORMATIVA"]), nc, pend,
                "OK" if d["_valid"] and not pend and not nc else "REVISAR"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = F_NORM
            if 2 <= j <= 6:
                c.number_format = NUM
        if vals[-1] != "OK":
            ws.cell(row=i, column=21).fill = FILL_ALERT
    # Pendências (texto integral da linha do PDF não lida) — abaixo da tabela.
    r = len(lista) + 4
    ws.cell(row=r, column=1, value="PENDÊNCIAS — linhas do PDF não lidas automaticamente "
            "(revisar manualmente; nenhuma foi somada)").font = F_SEC
    r += 1
    _hdr(ws, r, ["Tipo", "Competência", "Linha original do PDF"])
    r += 1
    tem = False
    for d in lista:
        for p in d["rubricas"].get("_PENDENCIAS", []):
            for j, v in enumerate(["PENDÊNCIA", d["meta"]["competencia"], p], 1):
                ws.cell(row=r, column=j, value=v).fill = FILL_ALERT
            r += 1
            tem = True
    if not tem:
        ws.cell(row=r, column=1, value="—")
        ws.cell(row=r, column=3, value="Nenhuma — todas as linhas de todos os PDFs foram lidas.")
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 13
    for j in range(2, 22):
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.column_dimensions["K"].width = 38
    ws.column_dimensions["O"].width = 30
    return ws


# ------------------------------------------------------------- Aba Resumo ---
def aba_resumo(wb, lista, categorias, enc_rows):
    ws = wb.create_sheet(SH_RES, 0)
    n = len(lista)
    cols = [get_column_letter(j) for j in range(2, n + 2)]
    ctot = get_column_letter(n + 2)
    last = cols[-1]

    ws["A1"] = "RESUMO POR ABA — Vencimentos, Descontos, Informativa e Encargos"
    ws["A1"].font = F_TITLE
    ws.cell(row=2, column=1, value="Categoria")
    for j, col in enumerate(cols):
        ws[f"{col}2"] = f"={SH_CON}!A{j + 2}"
    ws[f"{ctot}2"] = "Total período"
    _hdr(ws, 2, ["Categoria"] + [f"={SH_CON}!A{j + 2}" for j in range(n)] + ["Total período"])

    r = 3
    marks = {}

    def secao(txt):
        nonlocal r
        c = ws.cell(row=r, column=1, value=txt)
        c.font = F_SEC
        for j in range(1, n + 3):
            ws.cell(row=r, column=j).fill = FILL_SEC
        r += 1

    def linha_valores(rotulo, formula_fn, fill=None, total="SUM", font=F_NORM, fmt=NUM):
        nonlocal r
        c = ws.cell(row=r, column=1, value=rotulo)
        c.font = font
        for col in cols:
            cc = ws[f"{col}{r}"]
            cc.value, cc.number_format, cc.font = formula_fn(col), fmt, font
        tc = ws[f"{ctot}{r}"]
        tc.value = (f"=SUM(B{r}:{last}{r})" if total == "SUM" else
                    f"=AVERAGE(B{r}:{last}{r})" if total == "AVG" else total)
        tc.number_format, tc.font = fmt, F_BOLD
        if fill:
            for j in range(1, n + 3):
                ws.cell(row=r, column=j).fill = fill
        r += 1
        return r - 1

    def bloco_rubricas(titulo, aba, fixas, bloco, excluir_asterisco):
        nonlocal r
        secao(titulo)
        presentes = []
        for d in lista:
            for it in d["rubricas"].get(bloco, []):
                cat = categorias.get(it["codigo"], "NÃO CLASSIFICADO")
                if cat not in presentes:
                    presentes.append(cat)
        extras = [c for c in presentes if c not in fixas]  # nunca perder valor
        ini = r
        q = _q(aba)
        crit = f',{q}!$H:$H,"<>SIM"' if excluir_asterisco else ""
        for cat in fixas + extras:
            linha_valores(cat, lambda col, rr=r: (
                f"=SUMIFS({q}!$G:$G,{q}!$A:$A,{col}$2,{q}!$D:$D,$A{rr}{crit})"),
                fill=FILL_ALERT if cat in extras else None)
        fim = r - 1
        return ini, fim

    # 1. Vencimentos
    i, f = bloco_rubricas("1. VENCIMENTOS (aba Proventos)", SH_PRO, CATS_PROVENTOS, "PROVENTOS", True)
    marks["venc"] = linha_valores("TOTAL VENCIMENTOS", lambda col: f"=SUM({col}{i}:{col}{f})",
                                  fill=FILL_TOT, font=F_BOLD)
    r += 1
    # 2. Descontos
    i, f = bloco_rubricas("2. DESCONTOS (aba Descontos)", SH_DES, CATS_DESCONTOS, "DESCONTOS", True)
    marks["desc"] = linha_valores("TOTAL DESCONTOS", lambda col: f"=SUM({col}{i}:{col}{f})",
                                  fill=FILL_TOT, font=F_BOLD)
    r += 1
    # 3. Informativa
    i, f = bloco_rubricas("3. INFORMATIVA — não compõe líquido (aba Informativa)", SH_INF,
                          CATS_INFORMATIVA, "INFORMATIVA", False)
    marks["inf"] = linha_valores("TOTAL INFORMATIVA", lambda col: f"=SUM({col}{i}:{col}{f})",
                                 fill=FILL_TOT, font=F_BOLD)
    r += 1
    # 4. Encargos
    secao(f"4. ENCARGOS E IMPOSTOS PATRONAIS (aba {SH_ENC})")
    qe = _q(SH_ENC)
    i = r
    for rot, lab in ENCARGOS_RESUMO:
        er = enc_rows.get(lab)
        novo = lab in ("Empresa", "Contribuintes")
        linha_valores(rot, (lambda col, er=er: f"={qe}!{col}{er}") if er else (lambda col: 0),
                      fill=FILL_NEW if novo else (FILL_ALERT if not er else None))
    f = r - 1
    marks["enc"] = linha_valores("TOTAL ENCARGOS EMPRESA", lambda col: f"=SUM({col}{i}:{col}{f})",
                                 fill=FILL_TOT, font=F_BOLD)
    ws.cell(row=r, column=1, value=(
        "Linhas em verde: INSS patronal sobre folha (reoneração Lei 14.973/2024: 5% 2025, "
        "10% 2026, 15% 2027, 20% 2028). A CPRB sobre receita não consta do Resumo da Folha "
        "e não está neste custo.")).font = F_NOTE
    r += 2
    # 5. Fechamento
    secao("5. FECHAMENTO")
    v, dsc, e = marks["venc"], marks["desc"], marks["enc"]
    marks["liq"] = linha_valores("Líquido a pagar (Vencimentos - Descontos)",
                                 lambda col: f"={col}{v}-{col}{dsc}", font=F_BOLD)
    marks["custo"] = linha_valores("CUSTO TOTAL DA FOLHA (Vencimentos + Encargos)",
                                   lambda col: f"={col}{v}+{col}{e}", fill=FILL_KEY, font=F_BOLD)
    r += 1
    # 6. Headcount
    secao(f"6. HEADCOUNT (aba {SH_CON}, colunas G/H/I)")
    idx = lambda L: (lambda col: f"=INDEX({SH_CON}!${L}:${L},MATCH({col}$2,{SH_CON}!$A:$A,0))")
    a = linha_valores("Nº Trabalhando", idx("G"), total="AVG", fmt=INT)
    b = linha_valores("Admissões", idx("H"), fmt=INT)
    c = linha_valores("Demitidos", idx("I"), fmt=INT)
    linha_valores("Saldo do período (Admissões - Demitidos)", lambda col: f"={col}{b}-{col}{c}",
                  total=f"={ctot}{b}-{ctot}{c}", fmt=INT)
    r += 1
    # 7. Conferência contra o PDF (prova de que nada ficou de fora)
    secao("7. CONFERÊNCIA COM OS TOTAIS IMPRESSOS NO PDF (diferença deve ser zero)")
    for rot, L, ref in (("Vencimentos", "B", marks["venc"]), ("Descontos", "C", marks["desc"]),
                        ("Informativa", "D", marks["inf"]), ("Líquido Geral", "E", marks["liq"])):
        pr = linha_valores(f"{rot} — total impresso no PDF", idx(L))
        dr = linha_valores(f"{rot} — diferença (planilha - PDF)",
                           lambda col, pr=pr, ref=ref: f"=ROUND({col}{ref}-{col}{pr},2)")
        ws.conditional_formatting.add(
            f"B{dr}:{ctot}{dr}",
            CellIsRule(operator="notEqual", formula=["0"], fill=FILL_ALERT))

    for row in ws.iter_rows(min_row=3, max_row=r, max_col=n + 2):
        for cc in row:
            cc.border = Border(bottom=THIN)
    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 46
    for col in cols + [ctot]:
        ws.column_dimensions[col].width = 17
    return ws, marks


def preencher_custo_consolidado(wb, lista, marks):
    ws = wb[SH_CON]
    for i, col in enumerate(get_column_letter(j) for j in range(2, len(lista) + 2)):
        c = ws.cell(row=i + 2, column=6, value=f"={SH_RES}!{col}{marks['custo']}")
        c.number_format = NUM


# -------------------------------------------------------------- Validação ---
def validar(d):
    """Confere soma das linhas x Total impresso. Rubricas com '*' não compõem o
    Total impresso (comportamento do RM) — por isso ficam fora da soma."""
    msgs, ok = [], True
    for b in ("PROVENTOS", "DESCONTOS", "INFORMATIVA"):
        linhas = d["rubricas"].get(b, [])
        soma = sum(x["valor_calculado"] for x in linhas
                   if b == "INFORMATIVA" or not x["informativa"])
        imp = d["totais_impressos"].get(b)
        bate = imp is not None and abs(soma - imp) < 0.01
        ok &= bate
        msgs.append(f"{b}: linhas {soma:,.2f} x PDF {imp if imp is None else f'{imp:,.2f}'} "
                    f"{'OK' if bate else 'DIVERGE'}")
    t = d["totais_impressos"]
    if d["liquido_geral"] is not None and "PROVENTOS" in t and "DESCONTOS" in t:
        bate = abs(t["PROVENTOS"] - t["DESCONTOS"] - d["liquido_geral"]) < 0.01
        ok &= bate
        msgs.append(f"Líquido Geral {'OK' if bate else 'DIVERGE'}")
    for lab in ("Empresa", "Valor FGTS Rescisório", "Valor do FGTS", "RAT", "Terceiros"):
        if lab not in d["encargos"]:
            ok = False
            msgs.append(f"Encargo '{lab}' não encontrado")
    d["_valid"] = ok
    return msgs


def montar_excel(lista, categorias, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    aba_consolidado(wb, lista, categorias)
    aba_rubricas(wb, SH_PRO, lista, "PROVENTOS", categorias)
    aba_rubricas(wb, SH_DES, lista, "DESCONTOS", categorias)
    aba_rubricas(wb, SH_INF, lista, "INFORMATIVA", categorias)
    _, enc_rows = aba_encargos(wb, lista)
    _, marks = aba_resumo(wb, lista, categorias, enc_rows)
    preencher_custo_consolidado(wb, lista, marks)
    wb.active = 0
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("entradas", nargs="+", help="PDFs, pastas e/ou ZIPs de Resumo da Folha")
    ap.add_argument("--out", required=True)
    ap.add_argument("--categorias", default=str(Path(__file__).resolve().parent.parent
                                               / "assets" / "categorias_rubricas.csv"))
    args = ap.parse_args()

    categorias = carregar_categorias(args.categorias)
    with tempfile.TemporaryDirectory() as tmp:
        pdfs = coletar_pdfs(args.entradas, tmp)
        if not pdfs:
            sys.exit("ERRO: nenhum PDF encontrado nas entradas.")
        lista, ignorados = [], []
        for p in pdfs:
            d = parse_pdf(p)
            if not d["meta"].get("competencia") or d["totais_impressos"].get("PROVENTOS") is None:
                ignorados.append(p.name)  # não é um Resumo da Folha
                continue
            lista.append(d)

    if ignorados:
        print("IGNORADOS (não parecem 'RESUMO DA FOLHA'):", ", ".join(ignorados))
    if not lista:
        sys.exit("ERRO: nenhum PDF válido de Resumo da Folha.")

    empresas = {d["meta"].get("empresa") for d in lista}
    if len(empresas) > 1:
        sys.exit("ERRO: o lote mistura empresas (" + " | ".join(sorted(map(str, empresas))) +
                 "). Rode uma vez por empresa — somar empresas distintas no mesmo Resumo "
                 "distorce o custo.")
    comps = [d["meta"]["competencia"] for d in lista]
    dup = sorted({c for c in comps if comps.count(c) > 1})
    if dup:
        sys.exit(f"ERRO: competência repetida no lote ({', '.join(dup)}) — somaria em dobro. "
                 "Confira os arquivos (mesmo mês exportado duas vezes?).")

    lista.sort(key=lambda d: chave_comp(d["meta"]["competencia"]))
    for d in lista:
        msgs = validar(d)
        pend = len(d["rubricas"].get("_PENDENCIAS", []))
        nc = sorted({it["codigo"] for b in ("PROVENTOS", "DESCONTOS", "INFORMATIVA")
                     for it in d["rubricas"][b] if it["codigo"] not in categorias})
        print(f"[{d['meta']['competencia']}] {d['arquivo']} | Prov {len(d['rubricas']['PROVENTOS'])} "
              f"| Desc {len(d['rubricas']['DESCONTOS'])} | Inf {len(d['rubricas']['INFORMATIVA'])} "
              f"| Pendências {pend} | Não classificados {nc or 0}")
        print("    " + " | ".join(msgs))
    montar_excel(lista, categorias, args.out)
    print(f"OK -> {args.out}  ({len(lista)} competência(s))")


if __name__ == "__main__":
    main()
